"""Fetch NCES School Pulse Panel (SPP) staff-vacancy estimates.

Downloads the public "Staff Vacancies" workbook from NCES and writes
data/spp.json with, for every October survey wave in the file:
  - percent of positions vacant, by role (teaching and non-teaching), with SEs
  - percent of schools with at least one teaching vacancy
  - NCES's flags: "caution" / "unreliable" on individual estimates, and
    questionnaire_change on series whose questions changed without a bridge study

The SPP asked principals, for each role, the "total number of positions at my
school (filled + vacant)" and the "number of vacancies"; NCES reports vacancies
as a percent of positions, rounded to whole percentages.

Requires openpyxl (pip install openpyxl). If the download or parse fails, the
script exits non-zero and leaves the existing data/spp.json untouched.

Usage:  python scripts/fetch_spp.py
"""

import datetime as dt
import io
import json
import os
import re
import sys
import urllib.request

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "spp.json")
URL = "https://nces.ed.gov/surveys/spp/docs/release/StaffVac.xlsx"
UA = {"User-Agent": "Mozilla/5.0 (compatible; staff-shortages-dashboard/1.0)"}
# NCES footnote symbols, from the workbook's 'About' sheet.
FLAGS = {
    "!": "caution",          # coefficient of variation 30-50: interpret with caution
    "‡": "unreliable",  # reporting standards not met
}
QUESTIONNAIRE_CHANGE = "¹"  # questionnaire changed; no bridge study; compare time points with caution
MONTHS = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                                      "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}


def num(v):
    """SPP cells hold numbers, or symbols for 'not available', 'rounds to zero', etc."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.strip() == "#":
        return 0.0  # NCES: "Rounds to zero"
    return None


def wave_key(label):
    """'Oct24' -> '2024-10'."""
    m = re.fullmatch(r"([A-Z][a-z]{2})(\d{2})", label)
    return f"20{m.group(2)}-{MONTHS[m.group(1)]:02d}" if m else None


def main():
    req = urllib.request.Request(URL, headers=UA)
    with urllib.request.urlopen(req, timeout=120) as r:
        wb = openpyxl.load_workbook(io.BytesIO(r.read()), read_only=True, data_only=True)

    # Vacancy rates by role: sheet 'SVaverage', rows PCTVAC*_oct, all public schools.
    rows = list(wb["SVaverage"].iter_rows(values_only=True))
    hdr = rows[0]
    waves = [h[:-4] for h in hdr if isinstance(h, str) and h.endswith(" Pct")]
    roles = []
    for r in rows[1:]:
        if not r or r[1] != "All public schools" or not (r[2] or "").startswith("PCTVAC"):
            continue
        d = dict(zip(hdr, r))
        var = r[2]
        name = str(r[5]).split(":", 1)[-1].strip()
        group = "teaching" if var.startswith("PCTVAC1") else "non-teaching" if var.startswith("PCTVAC2") else "all"
        is_total = "TOT" in var
        values = {}
        for w in waves:
            k = wave_key(w)
            pct, se = num(d.get(f"{w} Pct")), num(d.get(f"{w} SE"))
            if k and pct is not None:
                v = {"pct": pct, "se": se}
                flag = FLAGS.get(str(d.get(f"{w} Flag") or "").strip())
                if flag:
                    v["flag"] = flag
                values[k] = v
        if values:
            roles.append({"var": var, "name": name, "group": group, "total": is_total,
                          # NCES marks series whose questionnaire changed with no bridge study.
                          "questionnaire_change": QUESTIONNAIRE_CHANGE in str(r[3] or ""),
                          "values": values})

    # Share of schools with >= 1 teaching vacancy: sheet 'All Schools', StaffVac1_TOT_CAT1.
    rows = list(wb["All Schools"].iter_rows(values_only=True))
    hdr = rows[0]
    waves2 = [h[:-4] for h in hdr if isinstance(h, str) and h.endswith(" Pct")]
    schools = {}
    for r in rows[1:]:
        if not r or r[2] != "StaffVac1_TOT_CAT1" or r[6] not in ("1 vacancy", ">1 vacancy"):
            continue
        d = dict(zip(hdr, r))
        for w in waves2:
            k, pct = wave_key(w), num(d.get(f"{w} Pct"))
            if k and pct is not None:
                schools.setdefault(k, 0.0)
                schools[k] += pct

    teach = next((x for x in roles if x["var"] == "PCTVAC1TOT_oct"), None)
    if not teach or len(teach["values"]) < 2:
        raise RuntimeError("SPP workbook layout changed: total teaching vacancy row not found")

    out = {
        "generated_utc": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%MZ"),
        "source_url": URL,
        "waves": sorted({k for x in roles for k in x["values"]}),
        "roles": roles,
        "schools_with_teaching_vacancy": dict(sorted(schools.items())),
    }
    tmp = OUT + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1)
    os.replace(tmp, OUT)
    print(f"OK: {len(roles)} role series; waves {out['waves']}; teaching = "
          + ", ".join(f"{k}: {v['pct']:g}%" for k, v in teach["values"].items()))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # noqa: BLE001
        print(f"FAILED, existing data left untouched: {e}", file=sys.stderr)
        sys.exit(1)
